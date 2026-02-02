import pandas as pd
import logging
import os

logger = logging.getLogger(__name__)

def export_backtest_results(results_map: dict, output_path: str):
    """
    백테스팅 결과를 엑셀로 저장하고 차트를 생성합니다.
    Args:
        results_map: { 'SheetName': DataFrame } 형태의 딕셔너리
        output_path: 저장할 파일 경로
    """
    try:
        import openpyxl
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("❌ 'openpyxl' 라이브러리가 필요합니다. (설치: python -m pip install openpyxl)")
        return

    try:
        # 디렉토리 생성
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            for sheet_name, df in results_map.items():
                if df.empty:
                    continue
                
                # 데이터 전처리 (차트용)
                df_export = df.copy()
                # 누적 수익금 계산
                if 'total_return' in df_export.columns:
                    df_export['cumulative_return'] = df_export['total_return'].cumsum()
                
                # 엑셀 저장
                df_export.to_excel(writer, sheet_name=sheet_name, index=False)
                
                # 워크북 및 시트 객체
                workbook = writer.book
                worksheet = writer.sheets[sheet_name]
                
                # 차트 생성 (데이터가 있고 필수 컬럼이 있을 때)
                if not df_export.empty and 'total_return' in df_export.columns and 'test_period' in df_export.columns:
                    # 컬럼 인덱스 찾기 (1-based index for openpyxl)
                    col_names = list(df_export.columns)
                    total_return_idx = col_names.index('total_return') + 1
                    test_period_idx = col_names.index('test_period') + 1
                    
                    # --- 차트 1: 기간별 손익 (Bar Chart) ---
                    bar_chart = BarChart()
                    bar_chart.type = "col"
                    bar_chart.style = 10
                    bar_chart.title = f"{sheet_name} - 기간별 손익"
                    bar_chart.y_axis.title = "손익 (KRW)"
                    bar_chart.x_axis.title = "기간"
                    bar_chart.height = 10
                    bar_chart.width = 20
                    
                    # 데이터 범위 설정 (헤더 포함)
                    data = Reference(worksheet, min_col=total_return_idx, min_row=1, max_row=len(df_export)+1)
                    cats = Reference(worksheet, min_col=test_period_idx, min_row=2, max_row=len(df_export)+1)
                    
                    bar_chart.add_data(data, titles_from_data=True)
                    bar_chart.set_categories(cats)
                    worksheet.add_chart(bar_chart, "J2")
                    
                    # --- 차트 2: 누적 수익금 (Line Chart) ---
                    if 'cumulative_return' in df_export.columns:
                        line_chart = LineChart()
                        line_chart.title = f"{sheet_name} - 누적 수익금"
                        line_chart.style = 13
                        line_chart.y_axis.title = "누적 수익 (KRW)"
                        line_chart.x_axis.title = "기간"
                        line_chart.height = 10
                        line_chart.width = 20
                        
                        cum_return_idx = col_names.index('cumulative_return') + 1
                        
                        data_cum = Reference(worksheet, min_col=cum_return_idx, min_row=1, max_row=len(df_export)+1)
                        line_chart.add_data(data_cum, titles_from_data=True)
                        line_chart.set_categories(cats)
                        
                        worksheet.add_chart(line_chart, "J22")
                
                # 컬럼 너비 자동 조정
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = get_column_letter(column[0].column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = (max_length + 2)
                    worksheet.column_dimensions[column_letter].width = adjusted_width

        logger.info(f"✅ 엑셀 변환 완료 (차트 포함): {output_path}")

    except Exception as e:
        logger.error(f"❌ 엑셀 저장 중 오류 발생: {e}")